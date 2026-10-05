import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    """
    Settea los bordes de una celda en python-docx.
    Uso: set_cell_border(cell, top={"sz": 4, "val": "single", "color": "000000"})
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        if border_name in kwargs:
            edge = OxmlElement(f'w:{border_name}')
            for key, val in kwargs[border_name].items():
                edge.set(qn(f'w:{key}'), str(val))
            tcBorders.append(edge)
    tcPr.append(tcBorders)

def make_header_doc(filename="encabezado_gimnasio.docx"):
    doc = Document()
    
    # Configurar márgenes de página
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Crear tabla de 2 filas y 3 columnas
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Combinar celdas de la primera columna (para el logo/escudo)
    table.cell(0, 0).merge(table.cell(1, 0))
    cell_logo = table.cell(0, 0)
    cell_title = table.cell(0, 1)
    cell_sub = table.cell(1, 1)
    cell_meta1 = table.cell(0, 2)
    cell_meta2 = table.cell(1, 2)

    # Anchos de columna
    widths = [Inches(1.2), Inches(4.5), Inches(1.8)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    # Estilo básico de borde negro fino
    border_kw = {
        'top': {'sz': 4, 'val': 'single', 'color': '000000'},
        'bottom': {'sz': 4, 'val': 'single', 'color': '000000'},
        'left': {'sz': 4, 'val': 'single', 'color': '000000'},
        'right': {'sz': 4, 'val': 'single', 'color': '000000'}
    }

    # Aplicar bordes y alineación vertical a todas las celdas
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, **border_kw)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # 1. Columna del Escudo / Logo
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_logo = p_logo.add_run("[ Insertar Escudo ]")
    r_logo.font.size = Pt(8)
    r_logo.font.italic = True
    r_logo.font.color.rgb = RGBColor(128, 128, 128)

    # 2. Celda Superior Central: Nombre y Lema
    p_title = cell_title.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_institution = p_title.add_run("GIMNASIO REAL AMERICANO\n")
    r_institution.bold = True
    r_institution.font.size = Pt(11)

    r_motto = p_title.add_run('"Formación axiológica con énfasis en emprendimiento desde la recreación,\nla cultura y el arte"')
    r_motto.font.size = Pt(8.5)
    r_motto.font.italic = True

    # 3. Celda Inferior Central (En blanco según solicitud)
    p_sub = cell_sub.paragraphs[0]
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("") # Dejado en blanco

    # 4. Celda Superior Derecha (En blanco según solicitud)
    p_meta1 = cell_meta1.paragraphs[0]
    p_meta1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_meta1.add_run("") # Dejado en blanco

    # 5. Celda Inferior Derecha (En blanco según solicitud)
    p_meta2 = cell_meta2.paragraphs[0]
    p_meta2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_meta2.add_run("") # Dejado en blanco

    doc.save(filename)
    print(f"Documento guardado exitosamente como: {filename}")

if __name__ == "__main__":
    make_header_doc()